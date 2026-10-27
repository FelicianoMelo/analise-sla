#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
gerar_total_sla.py
==================
Gera a planilha consolidada de SLA (Visão Geral / Total) idêntica à imagem de referência,
contendo os 3 painéis:
  1. TRANSPORTE DOMÉSTICO
  2. CLEARANCE – TRANSPORTE INTERNACIONAL
  3. SOLICITAÇÕES ( SCTASKs )

Os dados são lidos dinamicamente dos arquivos individuais gerados pelo script
'gerar_planilhas_sla.py' (*_DOM_aberto.xlsx, *_DOM_encerrado.xlsx, *_INT_aberto.xlsx,
*_INT_encerrado.xlsx, *_SCTASK_aberto.xlsx, *_SCTASK_encerrado.xlsx), NUNCA do arquivo
*_Total.xlsx anterior.
"""

from datetime import datetime, timedelta
import glob
import logging
import os
from pathlib import Path
import re
import sys
from typing import Dict, Any, Tuple, Optional

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import TextBlock, CellRichText
from openpyxl.cell.text import InlineFont
import pandas as pd

# ==============================================================================
# CONFIGURAÇÃO DE AMBIENTE E LOGGING
# ==============================================================================
PASTA_TRABALHO = Path(__file__).resolve().parent

logger = logging.getLogger("GERAR_TOTAL_SLA")
logger.setLevel(logging.INFO)

formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")

# Handler de arquivo
arquivo_log = logging.FileHandler(PASTA_TRABALHO / "analise_sla.log", encoding="utf-8")
arquivo_log.setFormatter(formatter)

# Handler de console
console = logging.StreamHandler(sys.stdout)
console.setFormatter(formatter)

if not logger.handlers:
    logger.addHandler(arquivo_log)
    logger.addHandler(console)


# ==============================================================================
# ESTILOS E PALETA DE CORES (IDÊNTICO À IMAGEM)
# ==============================================================================
# Preenchimentos dos Títulos dos Painéis
FILL_TITULO_DOM = PatternFill("solid", fgColor="C4D96D")      # Verde claro / lima
FILL_TITULO_INT = PatternFill("solid", fgColor="FADBD8")      # Rosa / salmão suave
FILL_TITULO_SCTASK = PatternFill("solid", fgColor="EAD1DC")   # Lavanda / lilás suave

# Subcabeçalhos (em Aberto, Encerrados)
FILL_SUBHEADER = PatternFill("solid", fgColor="D9D9D9")       # Cinza claro

# Tabelas Superiores
FILL_ABERTO_HEADER = PatternFill("solid", fgColor="F2C94C")   # Dourado / âmbar quente
FILL_ABERTO_TOTAL = PatternFill("solid", fgColor="F2C94C")    # Dourado / âmbar quente
FILL_ENCERRADO_HEADER = PatternFill("solid", fgColor="FFF2CC") # Amarelo pastel suave
FILL_ENCERRADO_TOTAL = PatternFill("solid", fgColor="FFF2CC")  # Amarelo pastel suave

# Barras de SLA e Totais
FILL_SLA_TOTAL = PatternFill("solid", fgColor="B7D7F0")       # Azul suave
FILL_SLA_OTIMO = PatternFill("solid", fgColor="98FB98")       # Verde suave (Dentro do SLA / SLA ótimo)
FILL_SLA_MEDIO_ATENCAO = PatternFill("solid", fgColor="D4AF37") # Mostarda / ouro escuro
FILL_SLA_MEDIO = PatternFill("solid", fgColor="FFEB9C")       # Amarelo
FILL_SLA_ATENCAO = PatternFill("solid", fgColor="F4A460")     # Laranja suave
FILL_SLA_FORA = PatternFill("solid", fgColor="FF0000")        # Vermelho vivo

# Caixa de Encerrados à Direita
FILL_ENCERRADOS_BOX = PatternFill("solid", fgColor="C27BA0")  # Malva / rosa antigo / ameixa

# Tipografia
FONT_NAME = "Arial"
FONT_TITULO_DOM = Font(name=FONT_NAME, size=11, bold=True, color="000000")
FONT_TITULO_INT = Font(name=FONT_NAME, size=11, bold=True, color="900000")
FONT_TITULO_SCTASK = Font(name=FONT_NAME, size=11, bold=True, color="4B0082")

FONT_SUBHEADER = Font(name=FONT_NAME, size=10, bold=True, color="000000")
FONT_HEADER_TABELA = Font(name=FONT_NAME, size=10, bold=True, color="000000")
FONT_DADOS = Font(name=FONT_NAME, size=10, bold=False, color="000000")
FONT_DADOS_BOLD = Font(name=FONT_NAME, size=10, bold=True, color="000000")
FONT_TOTAL_TABELA = Font(name=FONT_NAME, size=10, bold=True, color="000000")

FONT_SLA_TOTAL = Font(name=FONT_NAME, size=14, bold=True, color="000000")
FONT_SLA_BARRA = Font(name=FONT_NAME, size=10, bold=True, color="000000")
FONT_SLA_FORA = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF") # Texto branco no vermelho
FONT_SLA_DETALHE = Font(name=FONT_NAME, size=9, bold=False, color="000000")

FONT_ENCERRADOS_BOX = Font(name=FONT_NAME, size=14, bold=True, color="000000")
FONT_SCTASK_TOTAL = Font(name=FONT_NAME, size=14, bold=True, color="000000")
FONT_SCTASK_GRUPO = Font(name=FONT_NAME, size=9, bold=False, color="000000")
FONT_RODAPE = Font(name=FONT_NAME, size=9, bold=True, color="000000")

# Bordas
BORDA_THIN_BLACK = Side(style="thin", color="000000")
BORDER_CELULA = Border(
    left=BORDA_THIN_BLACK,
    right=BORDA_THIN_BLACK,
    top=BORDA_THIN_BLACK,
    bottom=BORDA_THIN_BLACK
)

COR_DIVISOR_PAINEL = "8EA9DB"
BORDA_DIVISOR_MEDIUM = Side(style="medium", color=COR_DIVISOR_PAINEL)


# ==============================================================================
# FUNÇÕES DE EXTRAÇÃO E PROCESSAMENTO DOS DADOS
# ==============================================================================
def classificar_sla(percentual: Any) -> str:
    """Classifica o percentual de SLA nas faixas definidas."""
    if pd.isna(percentual):
        return "fora"
    try:
        val = float(percentual)
    except (ValueError, TypeError):
        return "fora"

    if val < 50.00:
        return "otimo"
    elif val < 75.00:
        return "medio"
    elif val < 100.00:
        return "atencao"
    else:
        return "fora"


def formatar_data(prefixo: str) -> Tuple[str, str]:
    """Retorna datas formatadas DD/MM/YY para aberto e encerrado."""
    try:
        dt = datetime.strptime(prefixo, "%Y%m%d")
        dt_ontem = dt - timedelta(days=1)
        return dt.strftime("%d/%m/%y"), dt_ontem.strftime("%d/%m/%y")
    except Exception:
        # Fallback caso prefixo não seja YYYYMMDD
        hoje = datetime.now()
        ontem = hoje - timedelta(days=1)
        return hoje.strftime("%d/%m/%y"), ontem.strftime("%d/%m/%y")


def obter_prefixo_mais_recente(pasta: Path) -> str:
    """Encontra o prefixo de data mais recente entre os arquivos gerados."""
    hoje_str = datetime.now().strftime("%Y%m%d")
    if (pasta / f"{hoje_str}_DOM_aberto.xlsx").exists() or (pasta / f"{hoje_str}_INT_aberto.xlsx").exists():
        return hoje_str

    padrao = str(pasta / "*_DOM_aberto.xlsx")
    arquivos = glob.glob(padrao)
    prefixos = []
    for arq in arquivos:
        nome = Path(arq).name
        match = re.match(r"^(\d+)_DOM_aberto\.xlsx$", nome)
        if match:
            prefixos.append(match.group(1))

    if prefixos:
        prefixos.sort(reverse=True)
        return prefixos[0]

    return hoje_str


def carregar_dados_dom(pasta: Path, prefixo: str) -> Dict[str, Any]:
    """Extrai os dados de Transporte Doméstico a partir dos arquivos DOM."""
    arq_aberto = pasta / f"{prefixo}_DOM_aberto.xlsx"
    arq_encerrado = pasta / f"{prefixo}_DOM_encerrado.xlsx"

    atribuicoes = {}
    sla_otimo = 0
    sla_medio = 0
    sla_atencao = 0
    sla_fora = 0
    detalhes = {"otimo": {}, "medio": {}, "atencao": {}, "fora": {}}
    sla_total = 0

    if arq_aberto.exists():
        try:
            df = pd.read_excel(arq_aberto, sheet_name=0)
            if not df.empty:
                if "Atribuição a" in df.columns:
                    s_atrib = df["Atribuição a"].dropna().value_counts()
                    atribuicoes = s_atrib.to_dict()

                col_sla = next((c for c in df.columns if "% SLA Resolution" in str(c)), None)
                if col_sla:
                    df["cat_sla"] = df[col_sla].apply(classificar_sla)
                    sla_total = len(df)
                    for cat in ["otimo", "medio", "atencao", "fora"]:
                        sub = df[df["cat_sla"] == cat]
                        if not sub.empty:
                            if cat == "otimo":
                                sla_otimo = len(sub)
                            elif cat == "medio":
                                sla_medio = len(sub)
                            elif cat == "atencao":
                                sla_atencao = len(sub)
                            elif cat == "fora":
                                sla_fora = len(sub)

                            if "IC afetado" in sub.columns:
                                s_ic = sub["IC afetado"].dropna().value_counts()
                                detalhes[cat] = s_ic.to_dict()
        except Exception as e:
            logger.error(f"Erro ao ler {arq_aberto.name}: {e}")

    resolvidos = {}
    total_encerrados = 0
    if arq_encerrado.exists():
        try:
            df_enc = pd.read_excel(arq_encerrado, sheet_name=0)
            if not df_enc.empty:
                col_res = next((c for c in ["Resolvido por", "Atribuição a"] if c in df_enc.columns), None)
                if col_res:
                    resolvidos = df_enc[col_res].dropna().value_counts().to_dict()
                total_encerrados = len(df_enc)
        except Exception as e:
            logger.error(f"Erro ao ler {arq_encerrado.name}: {e}")

    return {
        "atribuicoes": atribuicoes,
        "total_atribuidos": sum(atribuicoes.values()),
        "resolvidos": resolvidos,
        "total_resolvidos": total_encerrados,
        "sla_total": sla_total if sla_total > 0 else (sla_otimo + sla_medio + sla_atencao + sla_fora),
        "sla_otimo": sla_otimo,
        "detalhe_otimo": detalhes["otimo"],
        "sla_medio": sla_medio,
        "detalhe_medio": detalhes["medio"],
        "sla_atencao": sla_atencao,
        "detalhe_atencao": detalhes["atencao"],
        "sla_fora": sla_fora,
        "detalhe_fora": detalhes["fora"],
    }


def carregar_dados_int(pasta: Path, prefixo: str) -> Dict[str, Any]:
    """Extrai os dados de Clearance - Transporte Internacional a partir dos arquivos INT."""
    arq_aberto = pasta / f"{prefixo}_INT_aberto.xlsx"
    if not arq_aberto.exists():
        arq_aberto = pasta / f"{prefixo}_CSATH_aberto.xlsx"

    arq_encerrado = pasta / f"{prefixo}_INT_encerrado.xlsx"
    if not arq_encerrado.exists():
        arq_encerrado = pasta / f"{prefixo}_CSATH_encerrado.xlsx"

    atribuicoes = {}
    sla_otimo = 0
    sla_medio = 0
    sla_atencao = 0
    sla_fora = 0
    detalhes = {"otimo": {}, "medio": {}, "atencao": {}, "fora": {}}
    sla_total = 0

    if arq_aberto.exists():
        try:
            df = pd.read_excel(arq_aberto, sheet_name=0)
            if not df.empty:
                if "Atribuição a" in df.columns:
                    s_atrib = df["Atribuição a"].dropna().value_counts()
                    atribuicoes = s_atrib.to_dict()

                col_sla = next((c for c in df.columns if "% SLA Resolution" in str(c)), None)
                if col_sla:
                    df["cat_sla"] = df[col_sla].apply(classificar_sla)
                    sla_total = len(df)
                    for cat in ["otimo", "medio", "atencao", "fora"]:
                        sub = df[df["cat_sla"] == cat]
                        if not sub.empty:
                            if cat == "otimo":
                                sla_otimo = len(sub)
                            elif cat == "medio":
                                sla_medio = len(sub)
                            elif cat == "atencao":
                                sla_atencao = len(sub)
                            elif cat == "fora":
                                sla_fora = len(sub)

                            if "IC afetado" in sub.columns:
                                s_ic = sub["IC afetado"].dropna().value_counts()
                                detalhes[cat] = s_ic.to_dict()
        except Exception as e:
            logger.error(f"Erro ao ler {arq_aberto.name}: {e}")

    resolvidos = {}
    total_encerrados = 0
    if arq_encerrado.exists():
        try:
            df_enc = pd.read_excel(arq_encerrado, sheet_name=0)
            if not df_enc.empty:
                col_res = next((c for c in ["Resolvido por", "Atribuição a"] if c in df_enc.columns), None)
                if col_res:
                    resolvidos = df_enc[col_res].dropna().value_counts().to_dict()
                total_encerrados = len(df_enc)
        except Exception as e:
            logger.error(f"Erro ao ler {arq_encerrado.name}: {e}")

    return {
        "atribuicoes": atribuicoes,
        "total_atribuidos": sum(atribuicoes.values()),
        "resolvidos": resolvidos,
        "total_resolvidos": total_encerrados,
        "sla_total": sla_total if sla_total > 0 else (sla_otimo + sla_medio + sla_atencao + sla_fora),
        "sla_otimo": sla_otimo,
        "detalhe_otimo": detalhes["otimo"],
        "sla_medio": sla_medio,
        "detalhe_medio": detalhes["medio"],
        "sla_atencao": sla_atencao,
        "detalhe_atencao": detalhes["atencao"],
        "sla_fora": sla_fora,
        "detalhe_fora": detalhes["fora"],
    }


def carregar_dados_sctask(pasta: Path, prefixo: str) -> Dict[str, Any]:
    """Extrai os dados de Solicitações (SCTASKs) a partir dos arquivos SCTASK."""
    arq_aberto = pasta / f"{prefixo}_SCTASK_aberto.xlsx"
    arq_encerrado = pasta / f"{prefixo}_SCTASK_encerrado.xlsx"

    atribuicoes = {}
    grupos = {}
    total_sctasks = 0

    if arq_aberto.exists():
        try:
            df = pd.read_excel(arq_aberto, sheet_name=0)
            if not df.empty:
                total_sctasks = len(df)
                if "Atribuição a" in df.columns:
                    s_atrib = df["Atribuição a"].dropna().value_counts()
                    atribuicoes = s_atrib.to_dict()

                if "Grupo de atribuição" in df.columns:
                    s_grp = df["Grupo de atribuição"].dropna().value_counts()
                    grupos = s_grp.to_dict()
        except Exception as e:
            logger.error(f"Erro ao ler {arq_aberto.name}: {e}")

    resolvidos = {}
    total_encerrados = 0
    if arq_encerrado.exists():
        try:
            df_enc = pd.read_excel(arq_encerrado, sheet_name=0)
            if not df_enc.empty:
                col_res = next((c for c in ["Resolvido por", "Atribuição a"] if c in df_enc.columns), None)
                if col_res:
                    resolvidos = df_enc[col_res].dropna().value_counts().to_dict()
                total_encerrados = len(df_enc)
        except Exception as e:
            logger.error(f"Erro ao ler {arq_encerrado.name}: {e}")

    return {
        "atribuicoes": atribuicoes,
        "total_atribuidos": sum(atribuicoes.values()),
        "resolvidos": resolvidos,
        "total_resolvidos": total_encerrados,
        "total_sctasks": total_sctasks,
        "grupos": grupos,
    }


def obter_dados_mock_imagem() -> Dict[str, Any]:
    """Retorna os dados exatos contidos na imagem de referência para teste/demonstração."""
    return {
        "prefixo": "20261022",
        "dom": {
            "atribuicoes": {},
            "total_atribuidos": 0,
            "resolvidos": {},
            "total_resolvidos": 0,
            "sla_total": 1,
            "sla_otimo": 0,
            "detalhe_otimo": {},
            "sla_medio": 0,
            "detalhe_medio": {},
            "sla_atencao": 0,
            "detalhe_atencao": {},
            "sla_fora": 1,
            "detalhe_fora": {"LAC-BR-3-GIS": 1},
        },
        "int": {
            "atribuicoes": {
                "Enoque Pinheiro": 21,
                "Rafael Miranda Ferreira": 15,
            },
            "total_atribuidos": 36,
            "resolvidos": {},
            "total_resolvidos": 0,
            "sla_total": 36,
            "sla_otimo": 5,
            "detalhe_otimo": {
                "LAC-BR-Clearance": 3,
                "LAC-BR-3-CSATH": 2,
            },
            "sla_medio": 0,
            "detalhe_medio": {},
            "sla_atencao": 1,
            "detalhe_atencao": {
                "LAC-BR-Clearance": 1,
            },
            "sla_fora": 30,
            "detalhe_fora": {
                "LAC-BR-3-CSATH": 19,
                "LAC-BR-Clearance": 9,
                "LAC-BR-3-ClearanceFormal": 2,
            },
        },
        "sctask": {
            "atribuicoes": {
                "Victorio Bormolini": 1,
                "Leandro Amorim": 1,
                "Diego Magalhaes": 1,
                "Fernando Junior": 1,
                "Duvanir Trigo": 8,
            },
            "total_atribuidos": 12,
            "resolvidos": {},
            "total_resolvidos": 0,
            "total_sctasks": 13,
            "grupos": {
                "FXE_Dev_LAC_BR_clearanceL3": 1,
                "FXE_Dev_LAC_BR_gisl3": 2,
                "FXE_Dev_LAC_BR_TMSL3": 1,
                "FXE_Dev_LAC_BR_MSAFCTEL3": 1,
                "( FXE_SA_LAC_FXESALACADMINL3 )": 8,
            },
        },
    }


# ==============================================================================
# FUNÇÕES DE RENDERIZAÇÃO NO EXCEL
# ==============================================================================
def formatar_faixa_mesclada(
    ws,
    row: int,
    col_ini: int,
    col_fim: int,
    valor: Any,
    fill: Optional[PatternFill] = None,
    font: Optional[Font] = None,
    alignment: Optional[Alignment] = None,
    border: Optional[Border] = BORDER_CELULA,
):
    """Escreve um valor e formata uma faixa mesclada de células."""
    if col_fim > col_ini:
        ws.merge_cells(start_row=row, start_column=col_ini, end_row=row, end_column=col_fim)

    cel_principal = ws.cell(row=row, column=col_ini)
    cel_principal.value = valor

    if font:
        cel_principal.font = font
    if alignment:
        cel_principal.alignment = alignment

    for c in range(col_ini, col_fim + 1):
        cel = ws.cell(row=row, column=c)
        if fill:
            cel.fill = fill
        if border:
            cel.border = border


def formatar_celula_individual(
    ws,
    row: int,
    col: int,
    valor: Any,
    fill: Optional[PatternFill] = None,
    font: Optional[Font] = None,
    alignment: Optional[Alignment] = None,
    border: Optional[Border] = BORDER_CELULA,
):
    """Formata uma única célula com valor, fill, font, alignment e border."""
    cel = ws.cell(row=row, column=col)
    cel.value = valor
    if font:
        cel.font = font
    if fill:
        cel.fill = fill
    if alignment:
        cel.alignment = alignment
    if border:
        cel.border = border


def desenhar_painel_cabecalho_tabelas(
    ws,
    col_ini: int,
    titulo: str,
    fill_titulo: PatternFill,
    font_titulo: Font,
    data_aberto_str: str,
    data_encerrado_str: str,
    atribuicoes: Dict[str, int],
    resolvidos: Dict[str, int],
    is_sctask: bool = False,
) -> int:
    """Desenha o título, subcabeçalhos e as tabelas de Aberto e Encerrados do painel."""
    c = col_ini

    # 1. TÍTULO DO PAINEL (Linha 2, col_ini a col_ini + 7)
    formatar_faixa_mesclada(
        ws,
        row=2,
        col_ini=c,
        col_fim=c + 7,
        valor=titulo,
        fill=fill_titulo,
        font=font_titulo,
        alignment=Alignment(horizontal="center", vertical="center"),
    )

    # 2. SUBCABEÇALHOS (Linha 4)
    # em Aberto (c a c+3)
    formatar_faixa_mesclada(
        ws,
        row=4,
        col_ini=c,
        col_fim=c + 3,
        valor=f"em Aberto:  {data_aberto_str}",
        fill=FILL_SUBHEADER,
        font=FONT_SUBHEADER,
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    # Encerrados (c+4 a c+7)
    formatar_faixa_mesclada(
        ws,
        row=4,
        col_ini=c + 4,
        col_fim=c + 7,
        valor=f"Encerrados:  {data_encerrado_str}",
        fill=FILL_SUBHEADER,
        font=FONT_SUBHEADER,
        alignment=Alignment(horizontal="center", vertical="center"),
    )

    # 3. CABEÇALHOS DAS TABELAS (Linha 6)
    # Aberto: Atribuído a | quantidade
    formatar_faixa_mesclada(
        ws,
        row=6,
        col_ini=c,
        col_fim=c + 2,
        valor="Atribuído a",
        fill=FILL_ABERTO_HEADER,
        font=FONT_HEADER_TABELA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=6,
        col=c + 3,
        valor="quantidade",
        fill=FILL_ABERTO_HEADER,
        font=FONT_HEADER_TABELA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )

    # Encerrados: Resolvido por | quantidade
    formatar_faixa_mesclada(
        ws,
        row=6,
        col_ini=c + 4,
        col_fim=c + 6,
        valor="Resolvido por",
        fill=FILL_ENCERRADO_HEADER,
        font=FONT_HEADER_TABELA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=6,
        col=c + 7,
        valor="quantidade",
        fill=FILL_ENCERRADO_HEADER,
        font=FONT_HEADER_TABELA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )

    # 4. LINHAS DE DADOS - ABERTO
    linha_atual_aberto = 7
    total_abertos = 0

    if not atribuicoes:
        # Se vazio, desenha uma linha em branco com quantidade 0 (igual transporte doméstico na imagem)
        formatar_faixa_mesclada(
            ws,
            row=linha_atual_aberto,
            col_ini=c,
            col_fim=c + 2,
            valor="",
            font=FONT_DADOS,
            alignment=Alignment(horizontal="left", vertical="center"),
        )
        formatar_celula_individual(
            ws,
            row=linha_atual_aberto,
            col=c + 3,
            valor=0,
            font=FONT_DADOS,
            alignment=Alignment(horizontal="right", vertical="center"),
        )
        linha_atual_aberto += 1
    else:
        for nome, qtd in atribuicoes.items():
            formatar_faixa_mesclada(
                ws,
                row=linha_atual_aberto,
                col_ini=c,
                col_fim=c + 2,
                valor=str(nome),
                font=FONT_DADOS,
                alignment=Alignment(horizontal="left", vertical="center"),
            )

            # Caso especial Duvanir Trigo na SCTASK (adiciona 1* em vermelho)
            if is_sctask and "Duvanir" in str(nome):
                rt_val = CellRichText(
                    f"{qtd} ",
                    TextBlock(InlineFont(color="FF0000", b=True), "1*")
                )
                formatar_celula_individual(
                    ws,
                    row=linha_atual_aberto,
                    col=c + 3,
                    valor=rt_val,
                    alignment=Alignment(horizontal="right", vertical="center"),
                )
            else:
                formatar_celula_individual(
                    ws,
                    row=linha_atual_aberto,
                    col=c + 3,
                    valor=int(qtd),
                    font=FONT_DADOS,
                    alignment=Alignment(horizontal="right", vertical="center"),
                )

            total_abertos += int(qtd)
            linha_atual_aberto += 1

    # Linha TOTAL - Aberto
    formatar_faixa_mesclada(
        ws,
        row=linha_atual_aberto,
        col_ini=c,
        col_fim=c + 2,
        valor="TOTAL",
        fill=FILL_ABERTO_TOTAL,
        font=FONT_TOTAL_TABELA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha_atual_aberto,
        col=c + 3,
        valor=total_abertos,
        fill=FILL_ABERTO_TOTAL,
        font=FONT_TOTAL_TABELA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha_atual_aberto += 1

    # 5. LINHAS DE DADOS - ENCERRADOS
    linha_atual_enc = 7
    total_enc = 0

    if not resolvidos:
        # Se vazio, vai direto para a linha TOTAL com 0
        formatar_faixa_mesclada(
            ws,
            row=linha_atual_enc,
            col_ini=c + 4,
            col_fim=c + 6,
            valor="TOTAL",
            fill=FILL_ENCERRADO_TOTAL,
            font=FONT_TOTAL_TABELA,
            alignment=Alignment(horizontal="left", vertical="center"),
        )
        formatar_celula_individual(
            ws,
            row=linha_atual_enc,
            col=c + 7,
            valor=0,
            fill=FILL_ENCERRADO_TOTAL,
            font=FONT_TOTAL_TABELA,
            alignment=Alignment(horizontal="right", vertical="center"),
        )
        linha_atual_enc += 1
    else:
        for nome, qtd in resolvidos.items():
            formatar_faixa_mesclada(
                ws,
                row=linha_atual_enc,
                col_ini=c + 4,
                col_fim=c + 6,
                valor=str(nome),
                font=FONT_DADOS,
                alignment=Alignment(horizontal="left", vertical="center"),
            )
            formatar_celula_individual(
                ws,
                row=linha_atual_enc,
                col=c + 7,
                valor=int(qtd),
                font=FONT_DADOS,
                alignment=Alignment(horizontal="right", vertical="center"),
            )
            total_enc += int(qtd)
            linha_atual_enc += 1

        formatar_faixa_mesclada(
            ws,
            row=linha_atual_enc,
            col_ini=c + 4,
            col_fim=c + 6,
            valor="TOTAL",
            fill=FILL_ENCERRADO_TOTAL,
            font=FONT_TOTAL_TABELA,
            alignment=Alignment(horizontal="left", vertical="center"),
        )
        formatar_celula_individual(
            ws,
            row=linha_atual_enc,
            col=c + 7,
            valor=total_enc,
            fill=FILL_ENCERRADO_TOTAL,
            font=FONT_TOTAL_TABELA,
            alignment=Alignment(horizontal="right", vertical="center"),
        )
        linha_atual_enc += 1

    return max(linha_atual_aberto, linha_atual_enc)


def desenhar_painel_sla_inferior(
    ws,
    col_ini: int,
    linha_inicio: int,
    dados_sla: Dict[str, Any],
    label_otimo: str = "Dentro do SLA:",
) -> int:
    """Desenha a seção inferior de SLA (SLA - TOTAL, Dentro do SLA, Médio, Atenção, Fora)."""
    c = col_ini
    linha = linha_inicio

    sla_total = dados_sla.get("sla_total", 0)
    sla_otimo = dados_sla.get("sla_otimo", 0)
    detalhe_otimo = dados_sla.get("detalhe_otimo", {})

    sla_medio = dados_sla.get("sla_medio", 0)
    detalhe_medio = dados_sla.get("detalhe_medio", {})

    sla_atencao = dados_sla.get("sla_atencao", 0)
    detalhe_atencao = dados_sla.get("detalhe_atencao", {})

    sla_medio_em_atencao = sla_medio + sla_atencao

    sla_fora = dados_sla.get("sla_fora", 0)
    detalhe_fora = dados_sla.get("detalhe_fora", {})

    # 1. SLA - TOTAL (Linha 14 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="SLA - TOTAL",
        fill=FILL_SLA_TOTAL,
        font=FONT_SLA_TOTAL,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_total,
        fill=FILL_SLA_TOTAL,
        font=FONT_SLA_TOTAL,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    # 2. DENTRO DO SLA / SLA ÓTIMO (Linha 15 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor=label_otimo,
        fill=FILL_SLA_OTIMO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_otimo,
        fill=FILL_SLA_OTIMO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    # Detalhes Dentro do SLA
    if detalhe_otimo:
        for fila, qtd in detalhe_otimo.items():
            formatar_faixa_mesclada(
                ws,
                row=linha,
                col_ini=c,
                col_fim=c + 2,
                valor=str(fila),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            formatar_celula_individual(
                ws,
                row=linha,
                col=c + 3,
                valor=int(qtd),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            linha += 1

    # Linha em branco de espaçamento se necessário para alinhamento com a linha 20
    while linha < linha_inicio + 6:
        linha += 1

    # 3. SLA MÉDIO / EM ATENÇÃO (Linha 20 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="SLA médio /em atenção",
        fill=FILL_SLA_MEDIO_ATENCAO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_medio_em_atencao,
        fill=FILL_SLA_MEDIO_ATENCAO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    # Sub-barra SLA Médio (Linha 21 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="SLA Médio",
        fill=FILL_SLA_MEDIO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_medio,
        fill=FILL_SLA_MEDIO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    if detalhe_medio:
        for fila, qtd in detalhe_medio.items():
            formatar_faixa_mesclada(
                ws,
                row=linha,
                col_ini=c,
                col_fim=c + 2,
                valor=str(fila),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            formatar_celula_individual(
                ws,
                row=linha,
                col=c + 3,
                valor=int(qtd),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            linha += 1

    # Sub-barra SLA em atenção (Linha 22 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="SLA em atenção",
        fill=FILL_SLA_ATENCAO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_atencao,
        fill=FILL_SLA_ATENCAO,
        font=FONT_SLA_BARRA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    if detalhe_atencao:
        for fila, qtd in detalhe_atencao.items():
            formatar_faixa_mesclada(
                ws,
                row=linha,
                col_ini=c,
                col_fim=c + 2,
                valor=str(fila),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            formatar_celula_individual(
                ws,
                row=linha,
                col=c + 3,
                valor=int(qtd),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            linha += 1

    # Linha em branco de espaçamento antes de Fora de SLA
    while linha < linha_inicio + 10:
        linha += 1

    # 4. FORA DE SLA (Linha 24 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="Fora de SLA",
        fill=FILL_SLA_FORA,
        font=FONT_SLA_FORA,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=sla_fora,
        fill=FILL_SLA_FORA,
        font=FONT_SLA_FORA,
        alignment=Alignment(horizontal="right", vertical="center"),
    )
    linha += 1

    if detalhe_fora:
        for fila, qtd in detalhe_fora.items():
            formatar_faixa_mesclada(
                ws,
                row=linha,
                col_ini=c,
                col_fim=c + 2,
                valor=str(fila),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            formatar_celula_individual(
                ws,
                row=linha,
                col=c + 3,
                valor=int(qtd),
                font=FONT_SLA_DETALHE,
                alignment=Alignment(horizontal="right", vertical="center"),
                border=None,
            )
            linha += 1

    return linha


def desenhar_painel_sctask_inferior(
    ws,
    col_ini: int,
    linha_inicio: int,
    total_sctasks: int,
    grupos: Dict[str, int],
) -> int:
    """Desenha a seção inferior de SCTASK (TOTAL geral e lista de grupos)."""
    c = col_ini
    linha = linha_inicio

    # 1. TOTAL (Linha 14 na imagem)
    formatar_faixa_mesclada(
        ws,
        row=linha,
        col_ini=c,
        col_fim=c + 2,
        valor="TOTAL",
        font=FONT_SCTASK_TOTAL,
        alignment=Alignment(horizontal="left", vertical="center"),
        border=None,
    )
    formatar_celula_individual(
        ws,
        row=linha,
        col=c + 3,
        valor=total_sctasks,
        font=FONT_SCTASK_TOTAL,
        alignment=Alignment(horizontal="right", vertical="center"),
        border=None,
    )
    linha += 2  # Pula uma linha conforme imagem

    # 2. Grupos de atribuição
    for grupo, qtd in grupos.items():
        nome_grp = str(grupo)
        if "FXESALACADMINL3" in nome_grp and not nome_grp.startswith("("):
            nome_grp = f"( {nome_grp} )"

        formatar_faixa_mesclada(
            ws,
            row=linha,
            col_ini=c,
            col_fim=c + 2,
            valor=nome_grp,
            font=FONT_SCTASK_GRUPO,
            alignment=Alignment(horizontal="right", vertical="center"),
            border=None,
        )
        formatar_celula_individual(
            ws,
            row=linha,
            col=c + 3,
            valor=int(qtd),
            font=FONT_SCTASK_GRUPO,
            alignment=Alignment(horizontal="right", vertical="center"),
            border=None,
        )
        linha += 1

    return linha


def desenhar_caixa_encerrados_lateral(
    ws,
    col_ini: int,
    linha_inicio: int,
    total_encerrados: int,
):
    """Desenha a caixa rosa/malva 'Encerrados: X' na coluna de encerrados."""
    c = col_ini + 4  # Colunas de encerrados
    formatar_faixa_mesclada(
        ws,
        row=linha_inicio,
        col_ini=c,
        col_fim=c + 2,
        valor="Encerrados:",
        fill=FILL_ENCERRADOS_BOX,
        font=FONT_ENCERRADOS_BOX,
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    formatar_celula_individual(
        ws,
        row=linha_inicio,
        col=c + 3,
        valor=total_encerrados,
        fill=FILL_ENCERRADOS_BOX,
        font=FONT_ENCERRADOS_BOX,
        alignment=Alignment(horizontal="right", vertical="center"),
    )


def aplicar_borda_perimetro_painel(ws, col_ini: int, col_fim: int, row_ini: int, row_fim: int):
    """Aplica uma borda estilizada de contorno médio em torno de um painel."""
    borda_med = Side(style="medium", color=COR_DIVISOR_PAINEL)

    for r in range(row_ini, row_fim + 1):
        # Borda esquerda
        cel_esq = ws.cell(row=r, column=col_ini)
        cel_esq.border = Border(
            left=borda_med,
            right=cel_esq.border.right if cel_esq.border else None,
            top=cel_esq.border.top if cel_esq.border else None,
            bottom=cel_esq.border.bottom if cel_esq.border else None,
        )

        # Borda direita
        cel_dir = ws.cell(row=r, column=col_fim)
        cel_dir.border = Border(
            left=cel_dir.border.left if cel_dir.border else None,
            right=borda_med,
            top=cel_dir.border.top if cel_dir.border else None,
            bottom=cel_dir.border.bottom if cel_dir.border else None,
        )

    for c in range(col_ini, col_fim + 1):
        # Borda superior
        cel_sup = ws.cell(row=row_ini, column=c)
        cel_sup.border = Border(
            left=cel_sup.border.left if cel_sup.border else None,
            right=cel_sup.border.right if cel_sup.border else None,
            top=borda_med,
            bottom=cel_sup.border.bottom if cel_sup.border else None,
        )

        # Borda inferior
        cel_inf = ws.cell(row=row_fim, column=c)
        cel_inf.border = Border(
            left=cel_inf.border.left if cel_inf.border else None,
            right=cel_inf.border.right if cel_inf.border else None,
            top=cel_inf.border.top if cel_inf.border else None,
            bottom=borda_med,
        )


# ==============================================================================
# FUNÇÃO PRINCIPAL DE CONSTRUÇÃO DO WORKBOOK
# ==============================================================================
def gerar_planilha_total(
    pasta: Path = PASTA_TRABALHO,
    prefixo: Optional[str] = None,
    dados_custom: Optional[Dict[str, Any]] = None,
    caminho_saida: Optional[Path] = None,
) -> Path:
    """Gera a planilha completa Total_SLA com os 3 painéis idênticos à imagem."""
    if prefixo is None:
        prefixo = obter_prefixo_mais_recente(pasta)

    logger.info(f"Processando dados consolidados para o prefixo: {prefixo}")

    # Carregar dados
    if dados_custom is not None:
        dados = dados_custom
        data_aberto_str, data_encerrado_str = formatar_data(dados.get("prefixo", prefixo))
    else:
        data_aberto_str, data_encerrado_str = formatar_data(prefixo)
        dados = {
            "dom": carregar_dados_dom(pasta, prefixo),
            "int": carregar_dados_int(pasta, prefixo),
            "sctask": carregar_dados_sctask(pasta, prefixo),
        }

    # Criar workbook openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visão Geral"
    ws.views.sheetView[0].showGridLines = True

    # Definir larguras de colunas
    larguras_padrao = {
        1: 3,   # A (margem)
        # Painel 1: B..I (cols 2..9)
        2: 9, 3: 9, 4: 9, 5: 9,
        6: 9, 7: 9, 8: 9, 9: 9,
        # Separador: J, K
        10: 2, 11: 2,
        # Painel 2: L..S (cols 12..19)
        12: 9, 13: 9, 14: 9, 15: 9,
        16: 9, 17: 9, 18: 9, 19: 9,
        # Separador: T, U
        20: 2, 21: 2,
        # Painel 3: V..AC (cols 22..29)
        22: 9, 23: 9, 24: 9, 25: 9,
        26: 9, 27: 9, 28: 9, 29: 9,
        30: 3,  # AD (margem)
    }
    for col_idx, largura in larguras_padrao.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    # Definir alturas de linha chave
    ws.row_dimensions[1].height = 15
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 8
    ws.row_dimensions[6].height = 20

    # 1. DESENHAR CABEÇALHOS E TABELAS SUPERIORES (LINHAS 2 A 12)
    fim_p1 = desenhar_painel_cabecalho_tabelas(
        ws,
        col_ini=2,
        titulo="TRANSPORTE DOMÉSTICO",
        fill_titulo=FILL_TITULO_DOM,
        font_titulo=FONT_TITULO_DOM,
        data_aberto_str=data_aberto_str,
        data_encerrado_str=data_encerrado_str,
        atribuicoes=dados["dom"]["atribuicoes"],
        resolvidos=dados["dom"]["resolvidos"],
        is_sctask=False,
    )

    fim_p2 = desenhar_painel_cabecalho_tabelas(
        ws,
        col_ini=12,
        titulo="CLEARANCE – TRANSPORTE INTERNACIONAL",
        fill_titulo=FILL_TITULO_INT,
        font_titulo=FONT_TITULO_INT,
        data_aberto_str=data_aberto_str,
        data_encerrado_str=data_encerrado_str,
        atribuicoes=dados["int"]["atribuicoes"],
        resolvidos=dados["int"]["resolvidos"],
        is_sctask=False,
    )

    fim_p3 = desenhar_painel_cabecalho_tabelas(
        ws,
        col_ini=22,
        titulo="SOLICITAÇÕES ( SCTASKs )",
        fill_titulo=FILL_TITULO_SCTASK,
        font_titulo=FONT_TITULO_SCTASK,
        data_aberto_str=data_aberto_str,
        data_encerrado_str=data_encerrado_str,
        atribuicoes=dados["sctask"]["atribuicoes"],
        resolvidos=dados["sctask"]["resolvidos"],
        is_sctask=True,
    )

    # A linha base inferior onde começa "SLA - TOTAL" é alinhada horizontalmente (linha 14 na imagem)
    linha_base_inferior = max(14, max(fim_p1, fim_p2, fim_p3) + 2)
    ws.row_dimensions[linha_base_inferior].height = 26

    # 2. DESENHAR SEÇÕES INFERIORES DE CADA PAINEL
    # Painel 1: Transporte Doméstico SLA
    fim_inf_p1 = desenhar_painel_sla_inferior(
        ws,
        col_ini=2,
        linha_inicio=linha_base_inferior,
        dados_sla=dados["dom"],
        label_otimo="SLA ótimo",
    )
    desenhar_caixa_encerrados_lateral(
        ws,
        col_ini=2,
        linha_inicio=linha_base_inferior,
        total_encerrados=dados["dom"]["total_resolvidos"],
    )

    # Painel 2: Clearance SLA
    fim_inf_p2 = desenhar_painel_sla_inferior(
        ws,
        col_ini=12,
        linha_inicio=linha_base_inferior,
        dados_sla=dados["int"],
        label_otimo="Dentro do SLA:",
    )
    desenhar_caixa_encerrados_lateral(
        ws,
        col_ini=12,
        linha_inicio=linha_base_inferior,
        total_encerrados=dados["int"]["total_resolvidos"],
    )

    # Painel 3: SCTASK Total e Grupos
    fim_inf_p3 = desenhar_painel_sctask_inferior(
        ws,
        col_ini=22,
        linha_inicio=linha_base_inferior,
        total_sctasks=dados["sctask"]["total_sctasks"],
        grupos=dados["sctask"]["grupos"],
    )
    desenhar_caixa_encerrados_lateral(
        ws,
        col_ini=22,
        linha_inicio=linha_base_inferior,
        total_encerrados=dados["sctask"]["total_resolvidos"],
    )

    linha_maxima = max(fim_inf_p1, fim_inf_p2, fim_inf_p3, 27)

    # 3. APLICAR BORDAS DE CONTORNO EM CADA PAINEL
    aplicar_borda_perimetro_painel(ws, col_ini=2, col_fim=9, row_ini=2, row_fim=linha_maxima)
    aplicar_borda_perimetro_painel(ws, col_ini=12, col_fim=19, row_ini=2, row_fim=linha_maxima)
    aplicar_borda_perimetro_painel(ws, col_ini=22, col_fim=29, row_ini=2, row_fim=linha_maxima)

    # 4. LINHA HORIZONTAL DIVISÓRIA INFERIOR E NOTA DE RODAPÉ
    linha_divisor = linha_maxima + 2
    ws.row_dimensions[linha_divisor].height = 6

    # Barra horizontal de destaque
    borda_divisor_h = Side(style="medium", color=COR_DIVISOR_PAINEL)
    for c in range(2, 30):
        cel = ws.cell(row=linha_divisor, column=c)
        cel.border = Border(bottom=borda_divisor_h)

    # Nota de rodapé (Linha 32 na imagem)
    linha_rodape = linha_divisor + 2
    texto_rodape = "*1 - para efeito de acompanhamento, incidente(s) / SCTASK(s) atribuido(s) a Duvanit Trigo ( Infraestrutura )"
    formatar_faixa_mesclada(
        ws,
        row=linha_rodape,
        col_ini=2,
        col_fim=29,
        valor=texto_rodape,
        font=FONT_RODAPE,
        alignment=Alignment(horizontal="left", vertical="center"),
        border=None,
    )
    ws.row_dimensions[linha_rodape].height = 20

    # 5. SALVAR ARQUIVO
    if caminho_saida is None:
        caminho_saida = pasta / f"{prefixo}_Total.xlsx"

    wb.save(caminho_saida)
    wb.close()
    logger.info(f"Planilha consolidada gerada com sucesso em: {caminho_saida}")
    return caminho_saida


# ==============================================================================
# EXECUÇÃO PRINCIPAL (CLI)
# ==============================================================================
def main():
    logger.info("Iniciando execução do script gerar_total_sla.py...")

    args = sys.argv[1:]
    prefixo = None
    dados_custom = None

    if "--imagem" in args or "--demo" in args or "--mock" in args:
        logger.info("Modo de demonstração (--imagem / --demo) ativado. Utilizando dados exatos da imagem de referência.")
        dados_custom = obter_dados_mock_imagem()
        prefixo = dados_custom["prefixo"]
    elif args:
        prefixo = args[0]
        logger.info(f"Prefixo recebido via linha de comando: {prefixo}")

    caminho_arquivo = gerar_planilha_total(
        pasta=PASTA_TRABALHO,
        prefixo=prefixo,
        dados_custom=dados_custom,
    )
    print(f"\n[SUCESSO] Planilha Total SLA gerada com sucesso:\n-> {caminho_arquivo}\n")


if __name__ == "__main__":
    main()
